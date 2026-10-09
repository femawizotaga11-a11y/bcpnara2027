# สร้างเทมเพลตฐานข้อมูล (xlsx) จาก tools/schema.json — หัวตาราง + Dropdown + นิยามแผน BCP/Config (ไม่มีข้อมูลสถานการณ์)
# ใช้: python tools/make_xlsx.py  →  database/BCP_Narathiwat_database.xlsx
import json, os, datetime
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

here = os.path.dirname(os.path.abspath(__file__))
db = json.load(open(os.path.join(here, 'schema.json'), encoding='utf8'))
wb = Workbook(); wb.remove(wb.active)
now = datetime.datetime.now().strftime('%Y-%m-%dT%H:%M:%S')
for tab in db['order']:
    d = db['tabs'][tab]; ws = wb.create_sheet(tab)
    keys = [c['k'] for c in d['cols']]; ws.append(keys)
    for i, c in enumerate(d['cols'], 1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1E3A8A'); cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(i)].width = 38 if c['t'] == 'area' else 16
        col = get_column_letter(i)
        if c['t'] == 'sel' and c.get('o'):
            dv = DataValidation(type='list', formula1='"' + ','.join(c['o']) + '"', allow_blank=True, showErrorMessage=True)
            ws.add_data_validation(dv); dv.add(f'{col}2:{col}1000')
    ws.freeze_panes = 'A2'
    if tab == 'BCP':
        for s in db['bcpSeed']:
            ws.append([f"BC-{s[0]:02d}", s[0], s[1], 'รอดำเนินการ', 0, '', '', '', s[2], now, 'setup'])
    if tab == 'Config':
        for s in db['configSeed']:
            ws.append([f"CF-{s[0].upper()}", s[0], '', s[1], now, 'setup'])
out = os.path.join(here, '..', 'database', 'BCP_Narathiwat_database.xlsx')
wb.save(out); print('saved', os.path.normpath(out), 'tabs:', len(db['order']))
